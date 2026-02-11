"""Excel report generation for payroll outputs."""

from __future__ import annotations

from pathlib import Path
from typing import Iterable, List, Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import (
    NACH_ACCOUNT_TYPE,
    NACH_CREDIT_NARRATION_PREFIX,
    NACH_DEBIT_NARRATION,
    OUTPUT_DIR,
)
from modules.statutory_calculator import get_pt_slab_label, split_employer_pf
from modules.wage_calculator import build_department_summary
from utils.date_utils import month_short_name, next_month_due_date

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin")
)
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")


def _currency(num: float) -> float:
    return round(float(num), 2)


def _output_dir(month: int, year: int, output_dir: Path | None = None) -> Path:
    base = output_dir or (OUTPUT_DIR / f"{year}_{month:02d}")
    base.mkdir(parents=True, exist_ok=True)
    return base


def _style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _style_table(ws, start_row: int, end_row: int, end_col: int) -> None:
    for r in range(start_row, end_row + 1):
        for c in range(1, end_col + 1):
            cell = ws.cell(r, c)
            cell.border = THIN_BORDER
            if r > start_row:
                cell.alignment = Alignment(vertical="center")


def _auto_fit_columns(ws) -> None:
    for col_cells in ws.columns:
        col_idx = col_cells[0].column
        length = max(len(str(cell.value or "")) for cell in col_cells)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(length + 2, 10), 45)


def _write_total_row(ws, row_index: int, total_label_col: int, number_cols: Sequence[int]) -> None:
    ws.cell(row_index, total_label_col, "TOTAL")
    ws.cell(row_index, total_label_col).font = Font(bold=True)
    ws.cell(row_index, total_label_col).fill = TOTAL_FILL
    for col in number_cols:
        letter = get_column_letter(col)
        ws.cell(row_index, col, f"=SUM({letter}2:{letter}{row_index-1})")
        ws.cell(row_index, col).font = Font(bold=True)
        ws.cell(row_index, col).fill = TOTAL_FILL


def generate_department_salary_summary(
    employee_wages: Iterable[dict], month: int, year: int, output_dir: Path | None = None
) -> Path:
    """Generate department-wise summary and employee details excel."""
    rows = list(employee_wages)
    dept_df = build_department_summary(rows)
    emp_df = pd.DataFrame(rows)

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"
    headers = [
        "Department",
        "Total Employees",
        "Total Days",
        "Total OT Hrs",
        "Total Gross",
        "Total Deductions",
        "Total Net Payable",
    ]
    ws_summary.append(headers)
    for _, item in dept_df.iterrows():
        ws_summary.append(
            [
                item["department"],
                int(item["total_employees"]),
                _currency(item["total_days"]),
                _currency(item["total_ot_hrs"]),
                _currency(item["total_gross"]),
                _currency(item["total_deductions"]),
                _currency(item["total_net_payable"]),
            ]
        )
    if ws_summary.max_row > 1:
        _write_total_row(ws_summary, ws_summary.max_row + 1, 1, [2, 3, 4, 5, 6, 7])
    _style_header(ws_summary, 1)
    _style_table(ws_summary, 1, ws_summary.max_row, 7)
    ws_summary.auto_filter.ref = f"A1:G{ws_summary.max_row}"
    ws_summary.freeze_panes = "A2"

    ws_emp = wb.create_sheet("Employee Details")
    if emp_df.empty:
        ws_emp.append(["No data"])
    else:
        cols = [
            "emp_code",
            "emp_name",
            "department",
            "designation",
            "present_days",
            "ot_hours",
            "basic_wages",
            "ot_amount",
            "gross_salary",
            "total_deductions",
            "net_payable",
        ]
        ws_emp.append(cols)
        for row in emp_df[cols].itertuples(index=False):
            ws_emp.append(list(row))
        _style_header(ws_emp, 1)
        _style_table(ws_emp, 1, ws_emp.max_row, len(cols))
        ws_emp.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws_emp.max_row}"
        ws_emp.freeze_panes = "A2"

    _auto_fit_columns(ws_summary)
    _auto_fit_columns(ws_emp)
    out_path = _output_dir(month, year, output_dir) / f"Dept_Salary_Summary_{month_short_name(month)}_{year}.xlsx"
    wb.save(out_path)
    return out_path


def generate_bank_payment_excel(employee_wages: Iterable[dict], month: int, year: int, output_dir: Path | None = None) -> Path:
    """Generate bank payment statement with NEFT format sheet."""
    rows = list(employee_wages)
    wb = Workbook()
    ws = wb.active
    ws.title = "Bank Payment"

    columns = [
        "S.No",
        "Emp Code",
        "Employee Name",
        "Department",
        "Grade",
        "Bank A/c No",
        "IFSC",
        "Bank Name",
        "Days",
        "OT Hrs",
        "Basic Wages",
        "OT Amount",
        "Gross Salary",
        "PF Ded",
        "ESIC Ded",
        "PT Ded",
        "Advance",
        "Other Ded",
        "Total Deductions",
        "Net Payable",
    ]
    ws.append(columns)
    for idx, row in enumerate(rows, start=1):
        ws.append(
            [
                idx,
                row.get("emp_code", ""),
                row.get("emp_name", ""),
                row.get("department", ""),
                row.get("grade", row.get("designation", "")),
                row.get("bank_account_no", ""),
                row.get("ifsc_code", ""),
                row.get("bank_name", ""),
                _currency(row.get("present_days", 0)),
                _currency(row.get("ot_hours", 0)),
                _currency(row.get("basic_wages", 0)),
                _currency(row.get("ot_amount", 0)),
                _currency(row.get("gross_salary", 0)),
                _currency(row.get("pf_employee", 0)),
                _currency(row.get("esic_employee", 0)),
                _currency(row.get("pt", 0)),
                _currency(row.get("advance", 0)),
                _currency(row.get("other_deduction", 0)),
                _currency(row.get("total_deductions", 0)),
                _currency(row.get("net_payable", 0)),
            ]
        )

    if ws.max_row > 1:
        _write_total_row(ws, ws.max_row + 1, 2, [9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20])
    _style_header(ws, 1)
    _style_table(ws, 1, ws.max_row, len(columns))

    for col in range(11, 21):
        for r in range(2, ws.max_row + 1):
            ws.cell(r, col).number_format = '#,##0.00'
    ws.conditional_formatting.add(
        f"K2:T{ws.max_row}",
        CellIsRule(operator="lessThan", formula=["0"], font=Font(color="9C0006")),
    )

    ws.auto_filter.ref = f"A1:T{ws.max_row}"
    ws.freeze_panes = "A2"

    neft = wb.create_sheet("NEFT_Upload")
    neft_cols = ["S.No", "Beneficiary Name", "Account Number", "IFSC", "Amount", "Narration"]
    neft.append(neft_cols)
    for idx, row in enumerate(rows, start=1):
        neft.append(
            [
                idx,
                row.get("emp_name", ""),
                row.get("bank_account_no", ""),
                row.get("ifsc_code", ""),
                _currency(row.get("net_payable", 0)),
                f"Salary {month_short_name(month)}-{year}",
            ]
        )
    _style_header(neft, 1)
    _style_table(neft, 1, neft.max_row, len(neft_cols))
    neft.auto_filter.ref = f"A1:F{neft.max_row}"
    neft.freeze_panes = "A2"
    for r in range(2, neft.max_row + 1):
        neft.cell(r, 5).number_format = '#,##0.00'

    nach = wb.create_sheet("NACH_Upload")
    nach_cols = [
        "S.No",
        "Beneficiary Name",
        "Account Number",
        "Account Type",
        "IFSC",
        "Amount",
        "Debit Narration",
        "Credit Narration",
        "Mobile",
        "Email",
    ]
    nach.append(nach_cols)
    for idx, row in enumerate(rows, start=1):
        nach.append(
            [
                idx,
                row.get("emp_name", ""),
                row.get("bank_account_no", ""),
                NACH_ACCOUNT_TYPE,
                row.get("ifsc_code", ""),
                _currency(row.get("net_payable", 0)),
                NACH_DEBIT_NARRATION,
                f"{NACH_CREDIT_NARRATION_PREFIX} {month_short_name(month)}-{year}",
                row.get("mobile", ""),
                row.get("email", ""),
            ]
        )
    _style_header(nach, 1)
    _style_table(nach, 1, nach.max_row, len(nach_cols))
    nach.auto_filter.ref = f"A1:J{nach.max_row}"
    nach.freeze_panes = "A2"
    for r in range(2, nach.max_row + 1):
        nach.cell(r, 6).number_format = '#,##0.00'

    _auto_fit_columns(ws)
    _auto_fit_columns(neft)
    _auto_fit_columns(nach)
    out_path = _output_dir(month, year, output_dir) / f"Bank_Payment_{month_short_name(month)}_{year}.xlsx"
    wb.save(out_path)
    return out_path


def generate_nach_upload_csv(employee_wages: Iterable[dict], month: int, year: int, output_dir: Path | None = None) -> Path:
    """Generate bank-friendly NACH/ACH CSV export."""
    rows = list(employee_wages)
    payload = []
    for idx, row in enumerate(rows, start=1):
        payload.append(
            {
                "SNo": idx,
                "BeneficiaryName": row.get("emp_name", ""),
                "AccountNumber": row.get("bank_account_no", ""),
                "AccountType": NACH_ACCOUNT_TYPE,
                "IFSC": row.get("ifsc_code", ""),
                "Amount": _currency(row.get("net_payable", 0)),
                "DebitNarration": NACH_DEBIT_NARRATION,
                "CreditNarration": f"{NACH_CREDIT_NARRATION_PREFIX} {month_short_name(month)}-{year}",
                "Mobile": row.get("mobile", ""),
                "Email": row.get("email", ""),
            }
        )
    out_path = _output_dir(month, year, output_dir) / f"NACH_Upload_{month_short_name(month)}_{year}.csv"
    pd.DataFrame(payload).to_csv(out_path, index=False)
    return out_path


def generate_pf_summary(employee_wages: Iterable[dict], month: int, year: int, output_dir: Path | None = None) -> Path:
    """Generate PF summary with employee and employer share details."""
    rows = list(employee_wages)
    wb = Workbook()
    ws = wb.active
    ws.title = "PF Summary"
    headers = [
        "S.No",
        "UAN",
        "Employee Name",
        "Father/Husband Name",
        "DOB",
        "DOJ",
        "Exit Date",
        "Basic Wages",
        "Employee PF (12%)",
        "Employer EPF (3.67%)",
        "Employer EPS (8.33%)",
        "Employer PF (13%)",
        "Gross PF",
    ]
    ws.append(headers)

    for idx, row in enumerate(rows, start=1):
        basic = _currency(row.get("pf_basic", row.get("basic_wages", 0)))
        emp_pf = _currency(row.get("pf_employee", 0))
        emp_epf, emp_eps = split_employer_pf(basic)
        emp_total = _currency(row.get("pf_employer", 0))
        ws.append(
            [
                idx,
                row.get("uan_no", ""),
                row.get("emp_name", ""),
                row.get("father_husband_name", ""),
                row.get("dob", ""),
                row.get("doj", ""),
                "",
                basic,
                emp_pf,
                emp_epf,
                emp_eps,
                emp_total,
                _currency(emp_pf + emp_total),
            ]
        )

    if ws.max_row > 1:
        _write_total_row(ws, ws.max_row + 1, 3, [8, 9, 10, 11, 12, 13])

    instruction_row = ws.max_row + 2
    ws.cell(instruction_row, 1, "ECR Upload Instructions:")
    ws.cell(instruction_row, 1).font = Font(bold=True)
    ws.cell(instruction_row + 1, 1, "1. Login to EPFO portal and go to ECR upload section.")
    ws.cell(instruction_row + 2, 1, "2. Upload ECR file generated from this payroll.")
    ws.cell(instruction_row + 3, 1, "3. Verify challan and pay by 15th of next month.")

    _style_header(ws, 1)
    _style_table(ws, 1, ws.max_row, len(headers))
    ws.auto_filter.ref = f"A1:M{ws.max_row}"
    ws.freeze_panes = "A2"
    for col in [8, 9, 10, 11, 12, 13]:
        for r in range(2, ws.max_row + 1):
            ws.cell(r, col).number_format = '#,##0.00'
    _auto_fit_columns(ws)

    out_path = _output_dir(month, year, output_dir) / f"PF_Summary_{month_short_name(month)}_{year}.xlsx"
    wb.save(out_path)
    return out_path


def generate_pf_ecr_file(employee_wages: Iterable[dict], month: int, year: int, output_dir: Path | None = None) -> Path:
    """Generate EPFO ECR-ready CSV file for statutory upload."""
    rows = list(employee_wages)
    payload = []
    for row in rows:
        basic = _currency(row.get("pf_basic", row.get("basic_wages", 0)))
        employee_pf = _currency(row.get("pf_employee", 0))
        employer_epf, employer_eps = split_employer_pf(basic)
        payload.append(
            {
                "UAN": row.get("uan_no", ""),
                "MemberName": row.get("emp_name", ""),
                "GrossWages": _currency(row.get("gross_salary", 0)),
                "EPFWages": basic,
                "EPSWages": basic,
                "EDLIWages": basic,
                "NCPDays": 0,
                "RefundOfAdvances": 0,
                "EPFContributionRemitted": employee_pf,
                "EPSContributionRemitted": employer_eps,
                "EPFContributionEmployerShare": employer_epf,
            }
        )
    out_path = _output_dir(month, year, output_dir) / f"PF_ECR_{month_short_name(month)}_{year}.csv"
    pd.DataFrame(payload).to_csv(out_path, index=False)
    return out_path


def generate_esic_summary(employee_wages: Iterable[dict], month: int, year: int, output_dir: Path | None = None) -> Path:
    """Generate ESIC summary with only eligible employees."""
    rows = [row for row in employee_wages if row.get("esic_applicable")]
    wb = Workbook()
    ws = wb.active
    ws.title = "ESIC Summary"
    headers = [
        "S.No",
        "ESIC No",
        "Employee Name",
        "Father/Husband Name",
        "Days Worked",
        "Gross Wages",
        "Employee ESIC (0.75%)",
        "Employer ESIC (3.25%)",
        "Total ESIC",
        "Remarks",
    ]
    ws.append(headers)
    for idx, row in enumerate(rows, start=1):
        emp = _currency(row.get("esic_employee", 0))
        er = _currency(row.get("esic_employer", 0))
        ws.append(
            [
                idx,
                row.get("esic_no", ""),
                row.get("emp_name", ""),
                row.get("father_husband_name", ""),
                _currency(row.get("present_days", 0)),
                _currency(row.get("gross_salary", 0)),
                emp,
                er,
                _currency(emp + er),
                "Applicable",
            ]
        )
    if ws.max_row > 1:
        _write_total_row(ws, ws.max_row + 1, 3, [5, 6, 7, 8, 9])
    inst = ws.max_row + 2
    ws.cell(inst, 1, "Challan Instructions: Pay on ESIC portal by 15th of next month.")
    ws.cell(inst, 1).font = Font(bold=True)

    _style_header(ws, 1)
    _style_table(ws, 1, ws.max_row, len(headers))
    ws.auto_filter.ref = f"A1:J{ws.max_row}"
    ws.freeze_panes = "A2"
    for col in [6, 7, 8, 9]:
        for r in range(2, ws.max_row + 1):
            ws.cell(r, col).number_format = '#,##0.00'
    _auto_fit_columns(ws)

    out_path = _output_dir(month, year, output_dir) / f"ESIC_Summary_{month_short_name(month)}_{year}.xlsx"
    wb.save(out_path)
    return out_path


def generate_pt_summary(employee_wages: Iterable[dict], month: int, year: int, output_dir: Path | None = None) -> Path:
    """Generate professional tax summary for all PT deducted rows."""
    rows = [row for row in employee_wages if float(row.get("pt", 0) or 0) > 0]
    wb = Workbook()
    ws = wb.active
    ws.title = "PT Summary"
    headers = ["S.No", "Employee Code", "Employee Name", "Gross Salary", "PT Slab", "PT Amount", "Remarks"]
    ws.append(headers)
    for idx, row in enumerate(rows, start=1):
        gross = _currency(row.get("gross_salary", 0))
        ws.append(
            [
                idx,
                row.get("emp_code", ""),
                row.get("emp_name", ""),
                gross,
                get_pt_slab_label(gross),
                _currency(row.get("pt", 0)),
                "Deducted",
            ]
        )
    if ws.max_row > 1:
        _write_total_row(ws, ws.max_row + 1, 3, [4, 6])

    due = next_month_due_date(21, month, year)
    ws.cell(ws.max_row + 2, 1, f"Payment Due Date: {due.isoformat()}")
    ws.cell(ws.max_row + 3, 1, "Generate challan in Maharashtra State Tax portal and pay before due date.")

    _style_header(ws, 1)
    _style_table(ws, 1, ws.max_row, len(headers))
    ws.auto_filter.ref = f"A1:G{ws.max_row}"
    ws.freeze_panes = "A2"
    for col in [4, 6]:
        for r in range(2, ws.max_row + 1):
            ws.cell(r, col).number_format = '#,##0.00'
    _auto_fit_columns(ws)

    out_path = _output_dir(month, year, output_dir) / f"PT_Summary_{month_short_name(month)}_{year}.xlsx"
    wb.save(out_path)
    return out_path


def generate_advance_register(
    advances: Iterable[dict], month: int, year: int, output_dir: Path | None = None
) -> Path:
    """Generate monthly advance register template or data sheet."""
    rows = list(advances)
    wb = Workbook()
    ws = wb.active
    ws.title = "Advance Register"
    headers = [
        "S.No",
        "Employee Name",
        "Father's Name",
        "Date of Advance",
        "Amount",
        "Purpose",
        "Installments",
        "Amount Recovered",
        "Balance",
        "Date of Recovery",
        "Signature",
    ]
    ws.append(headers)
    for idx, row in enumerate(rows, start=1):
        amount = _currency(row.get("amount", 0))
        recovered = _currency(row.get("amount_recovered", 0))
        ws.append(
            [
                idx,
                row.get("employee_name", ""),
                row.get("father_name", ""),
                row.get("date_of_advance", ""),
                amount,
                row.get("purpose", ""),
                row.get("installments", ""),
                recovered,
                _currency(amount - recovered),
                row.get("date_of_recovery", ""),
                row.get("signature", ""),
            ]
        )
    _style_header(ws, 1)
    _style_table(ws, 1, ws.max_row, len(headers))
    ws.auto_filter.ref = f"A1:K{ws.max_row}"
    ws.freeze_panes = "A2"
    for col in [5, 8, 9]:
        for r in range(2, ws.max_row + 1):
            ws.cell(r, col).number_format = '#,##0.00'
    _auto_fit_columns(ws)
    out_path = _output_dir(month, year, output_dir) / f"Advance_Register_{month_short_name(month)}_{year}.xlsx"
    wb.save(out_path)
    return out_path


def generate_wages_register_excel(
    employee_wages: Iterable[dict], month: int, year: int, output_dir: Path | None = None
) -> Path:
    """Generate wages register in Excel (Form II style data)."""
    rows = list(employee_wages)
    wb = Workbook()
    ws = wb.active
    ws.title = "Wages Register"
    headers = [
        "S.No",
        "Emp Code",
        "Employee Name",
        "Father/Husband Name",
        "Department",
        "Designation",
        "Days Worked",
        "OT Hrs",
        "Basic Wages",
        "OT Amount",
        "Gross Salary",
        "PF",
        "ESIC",
        "PT",
        "Advance",
        "Other Ded",
        "Total Deductions",
        "Net Wages Paid",
        "Payment Date",
        "Signature",
    ]
    ws.append(headers)
    for idx, row in enumerate(rows, start=1):
        ws.append(
            [
                idx,
                row.get("emp_code", ""),
                row.get("emp_name", ""),
                row.get("father_husband_name", ""),
                row.get("department", ""),
                row.get("designation", ""),
                _currency(row.get("present_days", 0)),
                _currency(row.get("ot_hours", 0)),
                _currency(row.get("basic_wages", 0)),
                _currency(row.get("ot_amount", 0)),
                _currency(row.get("gross_salary", 0)),
                _currency(row.get("pf_employee", 0)),
                _currency(row.get("esic_employee", 0)),
                _currency(row.get("pt", 0)),
                _currency(row.get("advance", 0)),
                _currency(row.get("other_deduction", 0)),
                _currency(row.get("total_deductions", 0)),
                _currency(row.get("net_payable", 0)),
                "",
                "",
            ]
        )
    if ws.max_row > 1:
        _write_total_row(ws, ws.max_row + 1, 3, [7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18])

    _style_header(ws, 1)
    _style_table(ws, 1, ws.max_row, len(headers))
    ws.auto_filter.ref = f"A1:T{ws.max_row}"
    ws.freeze_panes = "A2"
    for col in [9, 10, 11, 12, 13, 14, 15, 16, 17, 18]:
        for r in range(2, ws.max_row + 1):
            ws.cell(r, col).number_format = '#,##0.00'
    _auto_fit_columns(ws)
    out_path = _output_dir(month, year, output_dir) / f"Wages_Register_{month_short_name(month)}_{year}.xlsx"
    wb.save(out_path)
    return out_path

